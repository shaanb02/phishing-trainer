import csv
import io
import logging
import os
from datetime import datetime
from flask import render_template, request, redirect, url_for, flash, jsonify, make_response
from flask_login import login_user, logout_user, login_required, current_user
from app import app, db
from models import User, EmailTemplate, PhishingCampaign, SentEmail, ClickLog, AdminUser
from email_service import send_phishing_email
from ai_generator import generate_personalized_email
from tracking import generate_tracking_id, get_tracking_url


@app.route('/login', methods=['GET', 'POST'])
def login():
    if AdminUser.query.count() == 0:
        return redirect(url_for('setup'))
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        admin = AdminUser.query.filter_by(username=username, active=True).first()
        if admin and admin.check_password(password):
            admin.last_login = datetime.utcnow()
            db.session.commit()
            login_user(admin)
            next_page = request.args.get('next')
            return redirect(next_page or url_for('dashboard'))
        flash('Invalid username or password.', 'error')
    return render_template('login.html')


@app.route('/logout')
@login_required
def logout():
    logout_user()
    flash('You have been logged out.', 'info')
    return redirect(url_for('login'))


@app.route('/setup', methods=['GET', 'POST'])
def setup():
    if AdminUser.query.count() > 0:
        return redirect(url_for('login'))
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        email = request.form.get('email', '').strip()
        password = request.form.get('password', '')
        confirm = request.form.get('confirm_password', '')
        if not username or not email or not password:
            flash('All fields are required.', 'error')
        elif password != confirm:
            flash('Passwords do not match.', 'error')
        elif len(password) < 8:
            flash('Password must be at least 8 characters.', 'error')
        else:
            admin = AdminUser(username=username, email=email)
            admin.set_password(password)
            db.session.add(admin)
            db.session.commit()
            login_user(admin)
            flash('Admin account created. Welcome!', 'success')
            return redirect(url_for('dashboard'))
    return render_template('setup.html')


@app.route('/')
@login_required
def dashboard():
    """
    Main dashboard showing overview of phishing training campaigns
    Displays key metrics and recent activity
    """
    # Calculate overall statistics for the dashboard
    total_users = User.query.filter_by(active=True).count()
    total_campaigns = PhishingCampaign.query.count()
    active_campaigns = PhishingCampaign.query.filter_by(status='active').count()
    total_clicks = ClickLog.query.count()
    
    # Get recent campaigns for dashboard display
    recent_campaigns = PhishingCampaign.query.order_by(
        PhishingCampaign.created_at.desc()
    ).limit(5).all()
    
    # Get recent click activity for monitoring
    recent_clicks = ClickLog.query.join(User).order_by(
        ClickLog.clicked_at.desc()
    ).limit(10).all()
    
    # Calculate overall click rate across all campaigns
    total_sent = SentEmail.query.count()
    overall_click_rate = 0
    if total_sent > 0:
        overall_click_rate = round((total_clicks / total_sent) * 100, 2)
    
    return render_template('dashboard.html',
                         total_users=total_users,
                         total_campaigns=total_campaigns,
                         active_campaigns=active_campaigns,
                         total_clicks=total_clicks,
                         overall_click_rate=overall_click_rate,
                         recent_campaigns=recent_campaigns,
                         recent_clicks=recent_clicks)

@app.route('/users')
@login_required
def users():
    """
    User management page for viewing and managing employees/trainees
    Allows adding, editing, and viewing user information
    """
    users_list = User.query.filter_by(active=True).order_by(User.name).all()
    return render_template('users.html', users=users_list)

@app.route('/users/add', methods=['GET', 'POST'])
@login_required
def add_user():
    """
    Add a new user to the training system
    Handles form submission for creating new trainees
    """
    if request.method == 'POST':
        try:
            # Create new user from form data
            user = User(
                name=request.form['name'],
                email=request.form['email'],
                department=request.form['department'],
                role=request.form['role']
            )
            
            db.session.add(user)
            db.session.commit()
            
            flash(f'User {user.name} added successfully!', 'success')
            logging.info(f'New user added: {user.email}')
            
        except Exception as e:
            db.session.rollback()
            flash(f'Error adding user: {str(e)}', 'error')
            logging.error(f'Error adding user: {str(e)}')
        
        return redirect(url_for('users'))
    
    return render_template('users.html', show_add_form=True)

@app.route('/campaigns')
@login_required
def campaigns():
    """
    Campaign management page showing all phishing training campaigns
    Displays campaign status, statistics, and management options
    """
    campaigns_list = PhishingCampaign.query.order_by(
        PhishingCampaign.created_at.desc()
    ).all()
    
    templates = EmailTemplate.query.filter_by(active=True).all()
    
    return render_template('campaigns.html', 
                         campaigns=campaigns_list, 
                         templates=templates)

@app.route('/campaigns/<int:campaign_id>/details')
@login_required
def campaign_details(campaign_id):
    """Return campaign details with per-user send/click status as JSON."""
    campaign = PhishingCampaign.query.get_or_404(campaign_id)
    rows = []
    for se in campaign.emails_sent:
        click = se.clicks[0] if se.clicks else None
        rows.append({
            'name': se.user.name,
            'email': se.user.email,
            'department': se.user.department or '—',
            'role': se.user.role or '—',
            'sent_at': se.sent_at.strftime('%m/%d/%Y %H:%M') if se.sent_at else '—',
            'delivery_status': se.delivery_status,
            'clicked': bool(click),
            'clicked_at': click.clicked_at.strftime('%m/%d/%Y %H:%M') if click else None,
            'time_to_click': click.time_to_click if click else None,
            'device': click.device_type if click else None,
            'browser': click.browser if click else None,
        })
    return {
        'id': campaign.id,
        'name': campaign.name,
        'description': campaign.description or '',
        'status': campaign.status,
        'total_sent': campaign.total_sent,
        'total_clicked': campaign.total_clicked,
        'click_rate': campaign.click_rate,
        'started_at': campaign.started_at.strftime('%m/%d/%Y') if campaign.started_at else '—',
        'users': rows,
    }


@app.route('/campaigns/<int:campaign_id>/export')
@login_required
def export_campaign(campaign_id):
    """Export campaign recipient data as an Excel (.xlsx) file."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from flask import send_file

    campaign = PhishingCampaign.query.get_or_404(campaign_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Campaign Results'

    # ── Colour palette ──────────────────────────────────────────────────
    HEADER_FILL  = PatternFill('solid', fgColor='1F3864')   # dark navy
    CLICK_FILL   = PatternFill('solid', fgColor='FFCCCC')   # light red
    SAFE_FILL    = PatternFill('solid', fgColor='CCFFCC')   # light green
    TITLE_FONT   = Font(bold=True, color='FFFFFF', size=11)
    BOLD         = Font(bold=True)
    THIN         = Side(style='thin', color='CCCCCC')
    BORDER       = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    CENTER       = Alignment(horizontal='center', vertical='center')

    # ── Campaign summary block (rows 1-6) ───────────────────────────────
    summary = [
        ('Campaign Name',  campaign.name),
        ('Status',         campaign.status.title()),
        ('Started',        campaign.started_at.strftime('%m/%d/%Y') if campaign.started_at else '—'),
        ('Total Sent',     campaign.total_sent),
        ('Total Clicked',  campaign.total_clicked),
        ('Click Rate',     f'{campaign.click_rate}%'),
    ]
    for row_idx, (label, value) in enumerate(summary, start=1):
        ws.cell(row=row_idx, column=1, value=label).font = BOLD
        ws.cell(row=row_idx, column=2, value=value)

    # ── Column headers (row 8) ───────────────────────────────────────────
    headers = [
        'Name', 'Email', 'Department', 'Role',
        'Delivery Status', 'Sent At',
        'Clicked?', 'Clicked At', 'Time to Click (sec)',
        'Browser', 'Device',
    ]
    header_row = 8
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=h)
        cell.font = TITLE_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER

    # ── Data rows ────────────────────────────────────────────────────────
    for data_row, se in enumerate(campaign.emails_sent, start=header_row + 1):
        click = se.clicks[0] if se.clicks else None
        clicked = bool(click)
        row_fill = CLICK_FILL if clicked else SAFE_FILL

        values = [
            se.user.name,
            se.user.email,
            se.user.department or '',
            se.user.role or '',
            se.delivery_status,
            se.sent_at.strftime('%m/%d/%Y %H:%M') if se.sent_at else '',
            'Yes' if clicked else 'No',
            click.clicked_at.strftime('%m/%d/%Y %H:%M') if click else '',
            click.time_to_click if click else '',
            click.browser if click else '',
            click.device_type if click else '',
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=data_row, column=col, value=val)
            cell.fill = row_fill
            cell.border = BORDER
            cell.alignment = Alignment(vertical='center')

    # ── Column widths ────────────────────────────────────────────────────
    widths = [22, 30, 16, 16, 16, 18, 10, 18, 20, 16, 16]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

    ws.freeze_panes = f'A{header_row + 1}'   # freeze header row

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_name = ''.join(c if c.isalnum() or c in ' _-' else '_' for c in campaign.name)
    filename = f'{safe_name}_results.xlsx'

    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )


@app.route('/campaigns/create', methods=['POST'])
@login_required
def create_campaign():
    """
    Create a new phishing training campaign
    Can use pre-written templates or AI-generated content
    """
    try:
        # Create new campaign from form data
        campaign = PhishingCampaign(
            name=request.form['name'],
            description=request.form['description'],
            template_id=request.form.get('template_id') if request.form.get('template_id') else None,
            use_ai_generation=bool(request.form.get('use_ai')),
            ai_prompt_context=request.form.get('ai_context', ''),
            is_sequence='is_sequence' in request.form,
            max_emails=int(request.form.get('max_emails', 1)),
            interval_days=int(request.form.get('interval_days', 3))
        )
        
        db.session.add(campaign)
        db.session.commit()
        
        flash(f'Campaign "{campaign.name}" created successfully!', 'success')
        logging.info(f'New campaign created: {campaign.name}')
        
    except Exception as e:
        db.session.rollback()
        flash(f'Error creating campaign: {str(e)}', 'error')
        logging.error(f'Error creating campaign: {str(e)}')
    
    return redirect(url_for('campaigns'))

@app.route('/campaigns/<int:campaign_id>/launch', methods=['POST'])
@login_required
def launch_campaign(campaign_id):
    """
    Launch a phishing campaign by sending emails to all active users
    Handles both template-based and AI-generated emails with robust fallbacks
    """
    try:
        campaign = PhishingCampaign.query.get_or_404(campaign_id)
        
        if campaign.status == 'active':
            flash('Campaign is already active!', 'warning')
            return redirect(url_for('campaigns'))
        
        # Get all active users for the campaign
        users = User.query.filter_by(active=True).all()
        
        if not users:
            flash('No active users found to send emails to!', 'warning')
            return redirect(url_for('campaigns'))
        
        sent_count = 0
        failed_count = 0
        
        # Send emails to each user
        for user in users:
            try:
                subject = None
                html_content = None
                
                # Try AI generation if enabled
                if campaign.use_ai_generation:
                    try:
                        email_content = generate_personalized_email(
                            user_name=user.name,
                            user_role=user.role,
                            template_id=campaign.template_id,
                            context=campaign.ai_prompt_context
                        )
                        subject = email_content['subject']
                        html_content = email_content['html_content']
                        logging.info(f'AI generation successful for {user.email}')
                    except Exception as ai_err:
                        logging.warning(f'AI generation failed for {user.email}, falling back to template: {str(ai_err)}')
                
                # Fallback to template if AI failed or not enabled
                if not subject or not html_content:
                    template = campaign.template
                    if not template:
                        raise Exception("No template available for email content")
                    
                    subject = template.subject
                    html_content = template.html_content
                
                # Generate unique tracking ID for this email
                tracking_id = generate_tracking_id()
                
                # Replace tracking link placeholder in email content
                tracking_url = get_tracking_url(tracking_id)
                html_content = html_content.replace('{{TRACKING_URL}}', tracking_url)
                html_content = html_content.replace('{{USER_NAME}}', user.name)
                
                # Record the sent email in database
                sent_email = SentEmail(
                    campaign_id=campaign.id,
                    user_id=user.id,
                    subject=subject,
                    html_content=html_content,
                    tracking_id=tracking_id
                )
                
                db.session.add(sent_email)
                
                # Send the actual email
                send_success = send_phishing_email(
                    to_email=user.email,
                    subject=subject,
                    html_content=html_content
                )
                
                if send_success:
                    sent_email.delivery_status = 'sent'
                    sent_count += 1
                else:
                    sent_email.delivery_status = 'failed'
                    failed_count += 1
                
            except Exception as e:
                logging.error(f'Failed to process email for {user.email}: {str(e)}')
                failed_count += 1
        
        # Update campaign status and statistics
        campaign.status = 'active'
        campaign.started_at = datetime.utcnow()
        campaign.total_sent = sent_count
        
        db.session.commit()
        
        if sent_count > 0:
            flash(f'Campaign launched! Sent {sent_count} emails successfully.', 'success')
        if failed_count > 0:
            flash(f'Notice: {failed_count} emails failed to send. Check logs for details.', 'warning')
            
    except Exception as e:
        db.session.rollback()
        flash(f'Error launching campaign: {str(e)}', 'error')
        logging.error(f'Error launching campaign: {str(e)}')
    
    return redirect(url_for('campaigns'))

@app.route('/users/bulk-import', methods=['POST'])
@login_required
def bulk_import_users():
    """
    Bulk import users from a list of email addresses
    Accepts emails separated by commas, spaces, or newlines
    """
    import re
    try:
        raw_input = request.form.get('emails', '')
        # Split on commas, whitespace (spaces, tabs, newlines)
        emails = re.split(r'[\s,]+', raw_input.strip())
        emails = [e.strip().lower() for e in emails if e.strip()]

        added = 0
        skipped = 0
        invalid = 0

        for email in emails:
            # Basic email validation
            if not re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$', email):
                invalid += 1
                continue

            # Check if email already exists
            existing = User.query.filter_by(email=email).first()
            if existing:
                skipped += 1
                continue

            # Use the part before @ as the name
            name = email.split('@')[0].replace('.', ' ').replace('_', ' ').title()

            user = User(name=name, email=email)
            db.session.add(user)
            added += 1

        db.session.commit()

        parts = []
        if added:
            parts.append(f'{added} user{"s" if added != 1 else ""} added')
        if skipped:
            parts.append(f'{skipped} already existed')
        if invalid:
            parts.append(f'{invalid} invalid address{"es" if invalid != 1 else ""} skipped')

        flash(', '.join(parts) + '.', 'success' if added else 'warning')
        logging.info(f'Bulk import: {added} added, {skipped} skipped, {invalid} invalid')

    except Exception as e:
        db.session.rollback()
        flash(f'Error during bulk import: {str(e)}', 'error')
        logging.error(f'Bulk import error: {str(e)}')

    return redirect(url_for('users'))


@app.route('/templates')
@login_required
def templates():
    """
    Email template management page
    Shows all available phishing email templates
    """
    templates_list = EmailTemplate.query.filter_by(active=True).order_by(
        EmailTemplate.name
    ).all()
    
    return render_template('templates.html', templates=templates_list)

@app.route('/templates/add', methods=['POST'])
@login_required
def add_template():
    """
    Add a new email template to the system
    Templates can be used for campaigns or enhanced with AI
    """
    try:
        template = EmailTemplate(
            name=request.form['name'],
            subject=request.form['subject'],
            html_content=request.form['html_content'],
            text_content=request.form.get('text_content', ''),
            category=request.form['category'],
            difficulty_level=request.form['difficulty_level']
        )
        
        db.session.add(template)
        db.session.commit()
        
        flash(f'Template "{template.name}" added successfully!', 'success')
        logging.info(f'New template added: {template.name}')
        
    except Exception as e:
        db.session.rollback()
        flash(f'Error adding template: {str(e)}', 'error')
        logging.error(f'Error adding template: {str(e)}')
    
    return redirect(url_for('templates'))


@app.route('/templates/<int:template_id>/delete', methods=['POST'])
@login_required
def delete_template(template_id):
    """
    Delete a specific email template
    """
    try:
        template = EmailTemplate.query.get_or_404(template_id)
        template_name = template.name
        db.session.delete(template)
        db.session.commit()
        flash(f'Template "{template_name}" deleted successfully.', 'success')
        logging.info(f'Template deleted: {template_name}')
    except Exception as e:
        db.session.rollback()
        flash(f'Error deleting template: {str(e)}', 'error')
        logging.error(f'Error deleting template {template_id}: {str(e)}')
    return redirect(url_for('templates'))


@app.route('/campaigns/<int:campaign_id>/pause', methods=['POST'])
@login_required
def pause_campaign(campaign_id):
    """Pause a campaign so clicks are no longer tracked or recorded."""
    campaign = PhishingCampaign.query.get_or_404(campaign_id)
    campaign.status = 'paused'
    db.session.commit()
    flash(f'Campaign "{campaign.name}" has been paused. Tracking links are now disabled.', 'success')
    return redirect(url_for('campaigns'))


@app.route('/campaigns/<int:campaign_id>/resume', methods=['POST'])
@login_required
def resume_campaign(campaign_id):
    """Resume a paused campaign."""
    campaign = PhishingCampaign.query.get_or_404(campaign_id)
    campaign.status = 'active'
    db.session.commit()
    flash(f'Campaign "{campaign.name}" has been resumed. Tracking links are active again.', 'success')
    return redirect(url_for('campaigns'))


@app.route('/campaigns/<int:campaign_id>/delete', methods=['POST'])
@login_required
def delete_campaign(campaign_id):
    """
    Delete a campaign and all its associated sent emails and click logs
    """
    try:
        campaign = PhishingCampaign.query.get_or_404(campaign_id)
        campaign_name = campaign.name

        # Collect sent email IDs before any deletes
        sent_email_ids = [se.id for se in campaign.emails_sent]

        # Delete click logs first (deepest child), then sent emails, then campaign
        if sent_email_ids:
            ClickLog.query.filter(
                ClickLog.sent_email_id.in_(sent_email_ids)
            ).delete(synchronize_session='fetch')

        SentEmail.query.filter_by(
            campaign_id=campaign.id
        ).delete(synchronize_session='fetch')

        db.session.flush()   # push deletes before removing the parent row
        db.session.delete(campaign)
        db.session.commit()

        flash(f'Campaign "{campaign_name}" deleted successfully.', 'success')
        logging.info(f'Campaign deleted: {campaign_name}')

    except Exception as e:
        db.session.rollback()
        flash(f'Error deleting campaign: {str(e)}', 'error')
        logging.error(f'Error deleting campaign {campaign_id}: {str(e)}')

    return redirect(url_for('campaigns'))


@app.route('/users/clear-all', methods=['POST'])
@login_required
def clear_all_users():
    """
    Delete all users and their associated sent emails and click logs
    """
    try:
        # Delete click logs first, then sent emails, then users
        all_sent_email_ids = [se.id for se in SentEmail.query.all()]
        if all_sent_email_ids:
            ClickLog.query.filter(
                ClickLog.sent_email_id.in_(all_sent_email_ids)
            ).delete(synchronize_session='fetch')

        SentEmail.query.delete(synchronize_session='fetch')
        user_count = User.query.count()
        User.query.delete(synchronize_session='fetch')
        db.session.commit()

        flash(f'All {user_count} user(s) have been deleted successfully.', 'success')
        logging.info(f'All users cleared: {user_count} user(s) deleted.')

    except Exception as e:
        db.session.rollback()
        flash(f'Error clearing users: {str(e)}', 'error')
        logging.error(f'Error clearing all users: {str(e)}')

    return redirect(url_for('users'))


@app.route('/results')
@login_required
def results():
    """
    Results and analytics page showing campaign effectiveness
    Displays detailed statistics and allows CSV export
    """
    # Get all campaigns with their statistics
    campaigns = PhishingCampaign.query.filter(
        PhishingCampaign.status.in_(['active', 'completed'])
    ).all()
    
    # Get detailed click data for analysis
    click_data = db.session.query(
        ClickLog.clicked_at,
        ClickLog.time_to_click,
        ClickLog.device_type,
        ClickLog.browser,
        ClickLog.os,
        User.name,
        User.email,
        User.department,
        User.role,
        ClickLog.viewed_education,
        PhishingCampaign.name.label('campaign_name')
    ).join(
        User, ClickLog.user_id == User.id
    ).join(
        SentEmail, ClickLog.sent_email_id == SentEmail.id
    ).join(
        PhishingCampaign, SentEmail.campaign_id == PhishingCampaign.id
    ).order_by(ClickLog.clicked_at.desc()).all()
    
    return render_template('results.html', 
                         campaigns=campaigns, 
                         click_data=click_data)

@app.route('/template/send/<int:template_id>')
@login_required
def send_template(template_id):
    """
    Send a template directly to users without creating a full campaign
    Quick send functionality for testing templates
    """
    template = EmailTemplate.query.get_or_404(template_id)
    users = User.query.filter_by(active=True).all()
    
    return render_template('send_template.html', template=template, users=users)

@app.route('/template/send/<int:template_id>', methods=['POST'])
@login_required
def send_template_post(template_id):
    """
    Process sending template to selected users
    """
    template = EmailTemplate.query.get_or_404(template_id)
    
    try:
        # Get selected users or all users
        selected_user_ids = request.form.getlist('user_ids')
        send_to_all = 'send_to_all' in request.form
        
        if send_to_all:
            users = User.query.filter_by(active=True).all()
        else:
            users = User.query.filter(User.id.in_(selected_user_ids)).all()
        
        if not users:
            flash('No users selected for sending', 'warning')
            return redirect(url_for('send_template', template_id=template_id))
        
        # Server-side duplicate-send guard: block if this template was already
        # sent to any of these users within the last 5 minutes
        from datetime import timedelta
        cutoff = datetime.utcnow() - timedelta(minutes=5)
        recent_send = (
            db.session.query(SentEmail)
            .join(PhishingCampaign, SentEmail.campaign_id == PhishingCampaign.id)
            .filter(
                PhishingCampaign.template_id == template.id,
                SentEmail.user_id.in_([u.id for u in users]),
                SentEmail.sent_at >= cutoff
            )
            .first()
        )
        if recent_send:
            flash('This template was already sent to these users within the last 5 minutes. Please wait before sending again.', 'warning')
            return redirect(url_for('send_template', template_id=template_id))

        # Create a quick campaign for tracking
        campaign = PhishingCampaign(
            name=f'Quick Send: {template.name}',
            description=f'Direct template send - {template.name}',
            template_id=template.id,
            use_ai_generation=False,
            status='active',
            started_at=datetime.utcnow()
        )
        db.session.add(campaign)
        db.session.commit()
        
        sent_count = 0
        template_name = template.name  # Store template name for use outside loop
        
        for user in users:
            try:
                # Generate tracking ID
                tracking_id = generate_tracking_id()
                
                # Create personalized email content
                html_content = template.html_content.replace('{{USER_NAME}}', user.name)
                
                # Convert newlines to <br> tags to prevent "clump of text"
                html_content = html_content.replace('\n', '<br>')
                
                # Get base URL for tracking
                base_url = os.environ.get('REPLIT_DOMAINS', 'localhost:5000').split(',')[0]
                if not base_url.startswith('http'):
                    if 'replit' in base_url:
                        base_url = f'https://{base_url}'
                    else:
                        base_url = f'http://{base_url}'
                
                tracking_url = f'{base_url}/click/{tracking_id}'
                html_content = html_content.replace('{{TRACKING_URL}}', tracking_url)
                
                # Create sent email record
                sent_email = SentEmail(
                    campaign_id=campaign.id,
                    user_id=user.id,
                    subject=template.subject,
                    html_content=html_content,
                    tracking_id=tracking_id,
                    delivery_status='pending'
                )
                db.session.add(sent_email)
                
                # Send the email
                success = send_phishing_email(
                    to_email=user.email,
                    subject=template.subject,
                    html_content=html_content
                )
                
                if success:
                    sent_email.delivery_status = 'sent'
                    sent_count += 1
                else:
                    sent_email.delivery_status = 'failed'
                    
            except Exception as e:
                logging.error(f'Error sending to {user.email}: {str(e)}')
                continue
        
        # Update campaign totals
        campaign.total_sent = sent_count
        db.session.commit()
        
        if sent_count > 0:
            flash(f'Template sent successfully to {sent_count} users!', 'success')
            logging.info(f'Template {template_name} sent to {sent_count} users')
        else:
            flash('No emails were sent successfully', 'error')
            
    except Exception as e:
        db.session.rollback()
        flash(f'Error sending template: {str(e)}', 'error')
        logging.error(f'Error in send_template_post: {str(e)}')
    
    return redirect(url_for('templates'))

@app.route('/templates/content/<int:template_id>')
@login_required
def get_template_content(template_id):
    """
    API endpoint to fetch template content for preview
    Avoids JavaScript escaping issues in HTML
    """
    template = EmailTemplate.query.get_or_404(template_id)
    return jsonify({
        'html_content': template.html_content,
        'text_content': template.text_content
    })

@app.route('/results/export')
@login_required
def export_results():
    """
    Export click log results to CSV format
    Provides detailed data for further analysis and reporting
    """
    try:
        # Create CSV data in memory
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Write CSV headers
        writer.writerow([
            'Timestamp',
            'User Name',
            'Email',
            'Department',
            'Role',
            'Campaign',
            'IP Address',
            'User Agent',
            'Viewed Education'
        ])
        
        # Query and write click data
        clicks = db.session.query(
            ClickLog.clicked_at,
            User.name,
            User.email,
            User.department,
            User.role,
            PhishingCampaign.name.label('campaign_name'),
            ClickLog.ip_address,
            ClickLog.user_agent,
            ClickLog.viewed_education
        ).join(
            User, ClickLog.user_id == User.id
        ).join(
            SentEmail, ClickLog.sent_email_id == SentEmail.id
        ).join(
            PhishingCampaign, SentEmail.campaign_id == PhishingCampaign.id
        ).order_by(ClickLog.clicked_at.desc()).all()
        
        for click in clicks:
            writer.writerow([
                click.clicked_at.strftime('%Y-%m-%d %H:%M:%S'),
                click.name,
                click.email,
                click.department or '',
                click.role or '',
                click.campaign_name,
                click.ip_address or '',
                click.user_agent or '',
                'Yes' if click.viewed_education else 'No'
            ])
        
        # Create response with CSV file
        response = make_response(output.getvalue())
        response.headers['Content-Type'] = 'text/csv'
        response.headers['Content-Disposition'] = f'attachment; filename=phishing_results_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv'
        
        logging.info('Click log data exported to CSV')
        return response
        
    except Exception as e:
        flash(f'Error exporting data: {str(e)}', 'error')
        logging.error(f'Error exporting CSV: {str(e)}')
        return redirect(url_for('results'))

@app.route('/click/<tracking_id>')
def track_click(tracking_id):
    """
    Handle clicks on phishing links from emails
    This is the core tracking functionality that logs user interactions
    """
    try:
        # Find the sent email by tracking ID
        sent_email = SentEmail.query.filter_by(tracking_id=tracking_id).first()
        
        if not sent_email:
            logging.warning(f'Invalid tracking ID accessed: {tracking_id}')
            return render_template('phishing_landing.html', 
                                 error='Invalid tracking link'), 404

        # If the campaign is paused, the link is dead — show a neutral expired page
        campaign = sent_email.campaign
        if campaign and campaign.status == 'paused':
            logging.info(f'Click blocked for paused campaign "{campaign.name}" (tracking_id={tracking_id})')
            return render_template('phishing_landing.html',
                                 error='This training link is no longer active.'), 410

        # Check if this user already clicked (prevent duplicate logging)
        existing_click = ClickLog.query.filter_by(
            sent_email_id=sent_email.id,
            user_id=sent_email.user_id
        ).first()
        
        if not existing_click:
            # Parse user agent for advanced metrics
            from user_agents import parse
            ua_string = request.headers.get('User-Agent', '')
            ua = parse(ua_string)
            
            # Calculate time to click
            time_diff = (datetime.utcnow() - sent_email.sent_at).total_seconds()
            
            # Extract just the first (real client) IP from the proxy chain
            forwarded_for = request.environ.get('HTTP_X_FORWARDED_FOR', request.remote_addr)
            client_ip = forwarded_for.split(',')[0].strip()[:45]

            # Log the click with details for analysis
            click_log = ClickLog(
                sent_email_id=sent_email.id,
                user_id=sent_email.user_id,
                ip_address=client_ip,
                user_agent=ua_string,
                time_to_click=int(time_diff),
                device_type=ua.device.family,
                browser=ua.browser.family,
                os=ua.os.family
            )
            
            db.session.add(click_log)
            
            # Update campaign statistics
            campaign = sent_email.campaign
            campaign.total_clicked += 1
            
            db.session.commit()
            
            logging.info(f'Click logged for user {sent_email.user.email} in campaign {campaign.name}')
        
        # Show the educational landing page
        return render_template('phishing_landing.html',
                             user=sent_email.user,
                             campaign=sent_email.campaign,
                             tracking_id=tracking_id)
        
    except Exception as e:
        logging.error(f'Error tracking click for {tracking_id}: {str(e)}')
        return render_template('phishing_landing.html', 
                             error='An error occurred'), 500

@app.route('/education/<tracking_id>')
def show_education(tracking_id):
    """
    Show detailed cybersecurity education content
    Tracks when users view educational material after clicking
    """
    try:
        # Find the click log entry to mark education as viewed
        sent_email = SentEmail.query.filter_by(tracking_id=tracking_id).first()
        
        if sent_email:
            click_log = ClickLog.query.filter_by(
                sent_email_id=sent_email.id,
                user_id=sent_email.user_id
            ).first()
            
            if click_log and not click_log.viewed_education:
                click_log.viewed_education = True
                click_log.education_viewed_at = datetime.utcnow()
                db.session.commit()
                
                logging.info(f'Education viewed by user {sent_email.user.email}')
        
        return render_template('education.html', tracking_id=tracking_id)
        
    except Exception as e:
        logging.error(f'Error showing education for {tracking_id}: {str(e)}')
        return render_template('education.html', error='An error occurred')

@app.errorhandler(404)
def not_found_error(error):
    """Handle 404 errors with a user-friendly page"""
    return render_template('phishing_landing.html', 
                         error='Page not found'), 404

@app.errorhandler(500)
def internal_error(error):
    """Handle 500 errors and rollback database session"""
    db.session.rollback()
    return render_template('phishing_landing.html', 
                         error='Internal server error'), 500
